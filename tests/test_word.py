"""Word (.docx) reader: built from python-docx documents and hand-written Word XML."""
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image

docx = pytest.importorskip("docx")
from docx.oxml import OxmlElement  # noqa: E402
from docx.shared import Inches  # noqa: E402

from rag_ingest.ingest import run_ingest  # noqa: E402
from rag_ingest.readers import list_documents  # noqa: E402
from rag_ingest.readers import read_pages as read_any  # noqa: E402
from rag_ingest.readers.word import read_pages  # noqa: E402
from rag_ingest.readers.word.inline import (ImageMark, PAGE_BREAK, RENDERED_BREAK,  # noqa: E402
                                            inline_parts)
from rag_ingest.readers.word.numbering import Numbering, format_number  # noqa: E402
from rag_ingest.readers.word.tables import markdown_table  # noqa: E402
from rag_ingest.stores import VisionStore  # noqa: E402

NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
      'xmlns:v="urn:schemas-microsoft-com:vml" '
      'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"')


def xml(body: str):
    return etree.fromstring(body.replace("NS", NS, 1))


def png(folder: Path, name: str, size, color="white") -> Path:
    path = folder / name
    Image.new("RGB", size, color).save(path)
    return path


@pytest.fixture
def manual(tmp_path) -> Path:
    """A small Thai manual: heading, steps, bullets, screenshots, icon, table, page break."""
    big = png(tmp_path, "big.png", (800, 400))
    cell_shot = png(tmp_path, "cell.png", (600, 300), "lightblue")
    icon = png(tmp_path, "icon.png", (40, 40), "red")

    doc = docx.Document()
    doc.add_heading("คู่มือผู้จัดการสาขา", level=1)
    doc.add_paragraph("บทนำ: ระบบบันทึกอุบัติเหตุ")
    doc.add_paragraph("ขั้นแรก เข้าสู่ระบบ", style="List Number")
    doc.add_paragraph("ขั้นที่สอง กดบันทึก", style="List Number")
    doc.add_paragraph("หมายเหตุย่อย", style="List Bullet")
    doc.add_picture(str(big), width=Inches(4))
    doc.add_paragraph().add_run().add_picture(str(icon))
    table = doc.add_table(rows=3, cols=3)
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f"r{r}c{c}"
    table.cell(0, 0).merge(table.cell(0, 1)).text = "หัวตาราง|รวม"
    table.cell(1, 2).merge(table.cell(2, 2))
    table.cell(2, 0).paragraphs[0].add_run().add_picture(str(cell_shot), width=Inches(2))
    doc.add_page_break()
    doc.add_heading("ส่วนที่ 2 การแก้ไข", level=2)
    doc.add_picture(str(big), width=Inches(4))  # same screenshot again

    docs = tmp_path / "docs"
    docs.mkdir()
    path = docs / "manual.docx"
    doc.save(path)
    return path


def items(page):
    out = [(b.y, b.text) for b in page.blocks] + [(i.y, f"<{i.ref.stem}>") for i in page.images]
    return [text for _, text in sorted(out)]


def test_reading_order_headings_lists_tables_and_pages(manual):
    pages = list(read_pages(manual))
    assert [p.number for p in pages] == [1, 2]
    assert items(pages[0]) == [
        "# คู่มือผู้จัดการสาขา",
        "บทนำ: ระบบบันทึกอุบัติเหตุ",
        "1. ขั้นแรก เข้าสู่ระบบ",
        "2. ขั้นที่สอง กดบันทึก",
        "- หมายเหตุย่อย",
        "<manual_page_1_image_1>",
        "<manual_page_1_image_2>",
        "| หัวตาราง\\|รวม |  | r0c2 |\n|---|---|---|\n"
        "| r1c0 | r1c1 | r1c2<br>r2c2 |\n| r2c0 | r2c1 |  |",
        "<manual_page_1_image_3>",  # screenshot inside the table, after it
    ]
    assert items(pages[1]) == ["## ส่วนที่ 2 การแก้ไข", "<manual_page_2_image_1>"]
    assert pages[0].raw_text.startswith("# คู่มือผู้จัดการสาขา\nบทนำ")


def test_images_carry_pixels_size_and_bytes(manual):
    images = [i for p in read_pages(manual) for i in p.images]
    assert [(i.width, i.height, i.ext) for i in images] == [
        (800, 400, "png"), (40, 40, "png"), (600, 300, "png"), (800, 400, "png")]
    assert images[0].md5 == images[3].md5  # same screenshot -> OCR'd once by ingest


def test_ingest_end_to_end(manual, paths, tmp_path):
    store = VisionStore(paths.method("vision_typhoon"), paths)
    store.write("manual_page_1_image_1", {"image": "x", "text": "หน้าจอเข้าสู่ระบบ"})
    store.write("manual_page_1_image_3", {"image": "x", "text": "ตารางในหน้าจอ"})

    result = run_ingest(manual.parent, tmp_path / "out", [store])

    text = (tmp_path / "out" / "manual.md").read_text(encoding="utf-8")
    assert text.startswith("<!-- page 1 -->\n# คู่มือผู้จัดการสาขา\n\n")
    assert "2. ขั้นที่สอง กดบันทึก\n\n- หมายเหตุย่อย\n\n[ภาพหน้าจอ]\nหน้าจอเข้าสู่ระบบ" in text
    assert text.count("หน้าจอเข้าสู่ระบบ") == 2  # duplicate on page 2 reuses the OCR
    assert "<!-- page 2 -->\n## ส่วนที่ 2 การแก้ไข\n\n[ภาพหน้าจอ]\nหน้าจอเข้าสู่ระบบ" in text
    statuses = [(e["image"], e["status"]) for e in result.manifest]
    assert statuses == [("manual_page_1_image_1", "ocr"), ("manual_page_1_image_2", "icon"),
                        ("manual_page_1_image_3", "ocr"), ("manual_page_2_image_1", "duplicate")]
    assert result.manifest[0]["document"] == "manual.docx"
    assert result.manifest[0]["source"] == "vision_typhoon"


def test_rendered_page_breaks_win_over_hard_breaks(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("หน้าหนึ่ง")
    doc.add_page_break()  # ignored: the file has Word's own layout breaks
    later = doc.add_paragraph()
    later.add_run("ยังหน้าหนึ่ง")
    third = doc.add_paragraph()
    run = third.add_run()
    run._r.append(OxmlElement("w:lastRenderedPageBreak"))
    run.add_text("หน้าสอง")
    path = tmp_path / "rendered.docx"
    doc.save(path)

    pages = list(read_pages(path))
    assert [[b.text for b in p.blocks] for p in pages] == [["หน้าหนึ่ง", "ยังหน้าหนึ่ง"], ["หน้าสอง"]]


def test_text_boxes_alternate_content_and_tracked_changes():
    paragraph = xml("""<w:p NS>
      <w:r><w:t>ก่อน</w:t></w:r>
      <w:del><w:r><w:delText>ลบแล้ว</w:delText></w:r></w:del>
      <w:ins><w:r><w:t>เพิ่ม</w:t></w:r></w:ins>
      <w:r><w:instrText>PAGEREF _Toc1</w:instrText></w:r>
      <w:r><mc:AlternateContent>
        <mc:Choice><w:drawing><a:blip r:embed="rId5"/>
          <w:txbxContent><w:p><w:r><w:t>หมายเลข 1</w:t></w:r></w:p></w:txbxContent>
        </w:drawing></mc:Choice>
        <mc:Fallback><w:pict><v:imagedata r:id="rId5"/>
          <w:txbxContent><w:p><w:r><w:t>หมายเลข 1</w:t></w:r></w:p></w:txbxContent>
        </w:pict></mc:Fallback>
      </mc:AlternateContent></w:r>
      <w:r><w:tab/><w:t>หลัง</w:t><w:br/><w:br w:type="page"/><w:lastRenderedPageBreak/></w:r>
    </w:p>""")
    assert inline_parts(paragraph) == [
        "ก่อน", "เพิ่ม", ImageMark("rId5"), "\t", "หลัง", "\n", PAGE_BREAK, RENDERED_BREAK,
        "\n", "หมายเลข 1",  # text box once, after the paragraph
    ]


NUMBERING = """<w:numbering NS>
  <w:abstractNum w:abstractNumId="0">
    <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
    <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2"/></w:lvl>
  </w:abstractNum>
  <w:abstractNum w:abstractNumId="1">
    <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="thaiLetters"/><w:lvlText w:val="%1)"/></w:lvl>
  </w:abstractNum>
  <w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>
  <w:num w:numId="2"><w:abstractNumId w:val="1"/></w:num>
  <w:num w:numId="3"><w:abstractNumId w:val="0"/>
    <w:lvlOverride w:ilvl="0"><w:startOverride w:val="5"/></w:lvlOverride></w:num>
</w:numbering>"""


def test_numbering_levels_restart_and_formats():
    numbering = Numbering(xml(NUMBERING))
    labels = [numbering.label("1", 0), numbering.label("1", 1), numbering.label("1", 1),
              numbering.label("1", 0), numbering.label("1", 1)]
    assert labels == ["1.", "1.1", "1.2", "2.", "2.1"]  # sub-level restarts under a new item
    assert [numbering.label("2", 0) for _ in range(3)] == ["ก)", "ข)", "ค)"]
    assert numbering.label("3", 0) == "5."  # startOverride
    assert numbering.label("99", 0) is None  # unknown list


@pytest.mark.parametrize("n, fmt, expected", [
    (4, "decimal", "4"), (28, "lowerLetter", "bb"), (14, "upperRoman", "XIV"),
    (12, "thaiNumbers", "๑๒"), (3, "decimalZero", "03"), (2, "none", ""),
])
def test_format_number(n, fmt, expected):
    assert format_number(n, fmt) == expected


def test_markdown_table_drops_empty_rows_and_pads():
    assert markdown_table([["a", "b"], ["", ""], ["c"]]) == "| a | b |\n|---|---|\n| c |  |"
    assert markdown_table([]) == ""


def test_list_documents_rules(tmp_path, manual):
    folder = manual.parent
    (folder / "~$manual.docx").write_bytes(b"lock file")
    (folder / "notes.txt").write_text("x")
    assert list_documents(folder) == [manual]
    (folder / "manual.pdf").write_bytes(b"%PDF")
    with pytest.raises(ValueError, match="share a name"):
        list_documents(folder)
    with pytest.raises(ValueError, match="Unsupported document type"):
        read_any(folder / "notes.txt")
    empty = tmp_path / "empty"
    empty.mkdir()
    with pytest.raises(FileNotFoundError):
        list_documents(empty)


def test_list_documents_searches_type_subfolders(tmp_path, capsys):
    root = tmp_path / "raw"
    for sub in ("pdf", "docx", "txt", "md"):
        (root / sub).mkdir(parents=True)
        (root / sub / ".gitkeep").write_text("")
    (root / "pdf" / "b.pdf").write_bytes(b"%PDF")
    (root / "docx" / "a.docx").write_bytes(b"PK")
    (root / "docx" / "~$a.docx").write_bytes(b"lock")
    (root / "txt" / "notes.txt").write_text("x")
    (root / ".DS_Store").write_bytes(b"x")

    assert list_documents(root) == [root / "docx" / "a.docx", root / "pdf" / "b.pdf"]
    assert "Skipping 1 file(s) with no reader yet: txt/notes.txt" in capsys.readouterr().err

    (root / "docx" / "b.docx").write_bytes(b"PK")  # same name as pdf/b.pdf
    with pytest.raises(ValueError, match="share a name"):
        list_documents(root)
